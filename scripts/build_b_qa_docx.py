#!/usr/bin/env python3
"""Build a compact, template-organized DOCX from validated B-template records."""

from __future__ import annotations

import argparse
import json
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TEMPLATE_TITLES = {
    "B1": "Dangerous Frame Critique",
    "B2": "Is It Safe To Proceed?",
    "B3": "Risk Ranking",
    "B4": "Near-Miss Detection",
    "B5": "Stop Point",
    "B6": "Cause of Error",
    "B7": "CVS Safety Check",
}
TEMPLATE_ORDER = tuple(TEMPLATE_TITLES)

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "F3F5F7"
MID_GRAY = "D7DEE5"
TEXT = "222222"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text: str, bold: bool = False, color: str = TEXT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, text, end):
        run._r.append(node)


def add_bottom_border(paragraph, color: str = MID_GRAY, size: str = "6") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True

    for name, size, color, before, after in (
        ("Title", 28, NAVY, 0, 8),
        ("Subtitle", 14, DARK_BLUE, 0, 8),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 11.5, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "QA Label" not in styles:
        label = styles.add_style("QA Label", WD_STYLE_TYPE.CHARACTER)
    else:
        label = styles["QA Label"]
    label.font.name = "Calibri"
    label.font.size = Pt(10.5)
    label.font.bold = True
    label.font.color.rgb = RGBColor.from_string(NAVY)

    header = section.header
    header_para = header.paragraphs[0]
    header_para.text = "SURGICAL SAFETY QA  |  TEMPLATES B1–B7"
    header_para.style = styles["Normal"]
    header_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_para.paragraph_format.space_after = Pt(2)
    for run in header_para.runs:
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    add_bottom_border(header_para, color=BLUE, size="8")

    footer_para = section.footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_para.paragraph_format.space_after = Pt(0)
    lead = footer_para.add_run("SURGCOACH  •  ")
    lead.font.name = "Calibri"
    lead.font.size = Pt(8)
    lead.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    add_page_field(footer_para)
    for run in footer_para.runs[1:]:
        run.font.name = "Calibri"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(DARK_BLUE)


def add_labeled_paragraph(doc: Document, label: str, text: str, keep: bool = False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    if keep:
        keep_with_next(paragraph)
    paragraph.add_run(f"{label}: ").style = "QA Label"
    paragraph.add_run(text)
    return paragraph


def load_records(paths: list[Path]) -> list[dict]:
    records: list[dict] = []
    for path in paths:
        for line in path.read_text(encoding="utf-8").splitlines():
            if line.strip():
                record = json.loads(line)
                if record.get("validation_status") == "valid" and record.get("qa"):
                    records.append(record)
    order = {key: index for index, key in enumerate(TEMPLATE_ORDER)}
    return sorted(
        records,
        key=lambda row: (
            order.get(row["template_id"], 99),
            row.get("dataset", ""),
            row.get("video_id") or row.get("trial_id") or "",
            row.get("selection_id", ""),
        ),
    )


def build_document(records: list[dict], output: Path) -> None:
    counts = Counter(record["template_id"] for record in records)
    datasets: dict[str, Counter[str]] = defaultdict(Counter)
    for record in records:
        datasets[record["template_id"]][record.get("dataset", "Unknown")] += 1

    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(130)
    title.add_run("Surgical Safety QA")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Templates B1–B7 • Annotation-first generation")
    rule = doc.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rule.paragraph_format.space_before = Pt(10)
    rule.paragraph_format.space_after = Pt(18)
    rule.add_run("━" * 36).font.color.rgb = RGBColor.from_string(BLUE)
    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.add_run(
        f"{len(records)} validated question–answer pairs, organized by template"
    ).bold = True
    summary.paragraph_format.space_after = Pt(8)
    generated = doc.add_paragraph()
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER
    generated.add_run(date.today().strftime("%B %-d, %Y"))
    generated.runs[0].font.color.rgb = RGBColor.from_string(DARK_BLUE)

    doc.add_page_break()
    doc.add_heading("Corpus overview", level=1)
    overview = doc.add_paragraph(
        "Each entry uses the exact template question and reports video time in "
        "minutes and seconds. Model reasoning was required and verified during "
        "generation; private thinking traces are not reproduced here."
    )
    overview.paragraph_format.space_after = Pt(10)

    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    table.columns[0].width = Inches(1.05)
    table.columns[1].width = Inches(2.3)
    table.columns[2].width = Inches(3.75)
    hdr = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    for index, heading in enumerate(("Template", "Purpose", "Selected sources")):
        set_cell_text(hdr[index], heading, bold=True, color="FFFFFF")
        set_cell_shading(hdr[index], BLUE)
    for template_id in TEMPLATE_ORDER:
        cells = table.add_row().cells
        source_text = ", ".join(
            f"{dataset} ({count})"
            for dataset, count in sorted(datasets[template_id].items())
        )
        set_cell_text(cells[0], f"{template_id}  •  {counts[template_id]}", bold=True)
        set_cell_text(cells[1], TEMPLATE_TITLES[template_id])
        set_cell_text(cells[2], source_text)
        if len(table.rows) % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, LIGHT_GRAY)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    note.add_run("Selection note. ").bold = True
    note.add_run(
        "B1, B3, and B5 use three distinct videos from each of BernBypass70, "
        "StrasBypass70, and Endoscapes2023. B4 uses six deliberately difficult "
        "negative controls from the two bypass cohorts: the source labels "
        "document actual adverse events, so they must not be mislabeled as near "
        "misses. CholecT50 was not used for B1/B3/B4/B5 because its local "
        "annotations do not identify safety events."
    )

    for template_id in TEMPLATE_ORDER:
        template_heading = doc.add_heading(
            f"{template_id}  |  {TEMPLATE_TITLES[template_id]}",
            level=1,
        )
        template_heading.paragraph_format.page_break_before = True
        template_records = [
            record for record in records if record["template_id"] == template_id
        ]
        dataset_summary = ", ".join(
            f"{dataset}: {count}"
            for dataset, count in sorted(datasets[template_id].items())
        )
        intro = doc.add_paragraph(
            f"{len(template_records)} QA pairs • {dataset_summary}"
        )
        intro.paragraph_format.space_after = Pt(12)
        intro.runs[0].font.color.rgb = RGBColor.from_string(DARK_BLUE)
        intro.runs[0].italic = True

        for index, record in enumerate(template_records, start=1):
            item_id = (
                record.get("video_id")
                or record.get("trial_id")
                or record.get("selection_id")
                or "Unknown"
            )
            heading = doc.add_heading(
                f"{template_id}.{index}  •  {record.get('dataset', 'Unknown')}  "
                f"•  {item_id}  •  {record.get('timestamp_or_frame', '')}",
                level=3,
            )
            add_bottom_border(heading, color=MID_GRAY, size="4")
            qa = record["qa"][0]
            add_labeled_paragraph(doc, "Question", qa["question"], keep=True)
            add_labeled_paragraph(doc, "Answer", qa["answer"])
            rationale = add_labeled_paragraph(doc, "Rationale", qa["rationale"])
            rationale.paragraph_format.space_after = Pt(12)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("records", nargs="+", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    records = load_records(args.records)
    if not records:
        parser.error("No valid records found")
    build_document(records, args.output)
    print(f"wrote {args.output} ({len(records)} validated QA pairs)")


if __name__ == "__main__":
    main()
